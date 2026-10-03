"""对表格附件运行受限查询，并计算不含代码执行的数值表达式。"""

from __future__ import annotations

import ast
import csv
import hashlib
import json
import re
import sqlite3
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_FILE_BYTES = 30 * 1024 * 1024
MAX_EXPANDED_BYTES = 120 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 100
MAX_CELLS = 300_000
MAX_RESULT_ROWS = 500
MAX_RESULT_CHARS = 100_000
MAX_SQL_STEPS = 2_000_000
MAX_EXPRESSION_NODES = 80
MAX_MAGNITUDE = Decimal("1e100")
_SAFE_FUNCTIONS = {"abs", "avg", "coalesce", "count", "ifnull", "instr", "length", "lower", "ltrim", "max", "min", "nullif", "replace", "round", "rtrim", "substr", "sum", "total", "trim", "upper"}


def _fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _value(value: object) -> str | int | float | None:
    if value is None or isinstance(value, (str, int, float)):
        return value
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


def _color(color: object) -> str | None:
    kind = getattr(color, "type", None)
    value = getattr(color, kind, None) if kind in {"rgb", "theme", "indexed"} else None
    return str(value) if value is not None else None


def _new_table(connection: sqlite3.Connection, name: str, width: int) -> None:
    columns = ", ".join(f'"c{index}"' for index in range(1, width + 1))
    connection.execute(f'CREATE TABLE "{name}" (row_number INTEGER PRIMARY KEY, {columns})')


def _load_csv(connection: sqlite3.Connection, path: Path) -> list[dict[str, object]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.reader(source, delimiter=delimiter)
        width = 0
        rows = []
        for row in reader:
            if len(rows) >= MAX_ROWS or sum(map(len, row)) > MAX_RESULT_CHARS:
                raise ValueError("Table exceeds the supported row or field size limit")
            width = max(width, len(row))
            if width > MAX_COLUMNS or (len(rows) + 1) * width > MAX_CELLS:
                raise ValueError("Table exceeds the supported column or cell limit")
            rows.append(row)
    width = max(width, 1)
    _new_table(connection, "sheet_1", width)
    slots = ", ".join("?" for _ in range(width + 1))
    connection.executemany(
        f'INSERT INTO "sheet_1" VALUES ({slots})',
        ((index, *(row + [None] * (width - len(row)))) for index, row in enumerate(rows, 1)),
    )
    return [{"table": "sheet_1", "sheet": path.name, "rows": len(rows), "columns": ["row_number"] + [f"c{i}" for i in range(1, width + 1)]}]


def _load_xlsx(connection: sqlite3.Connection, path: Path) -> list[dict[str, object]]:
    with ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 5000 or sum(member.file_size for member in members) > MAX_EXPANDED_BYTES:
            raise ValueError("Workbook exceeds the supported expanded size limit")
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables = []
    total_cells = 0
    try:
        for index, sheet in enumerate(workbook.worksheets, 1):
            if sheet.max_row is None or sheet.max_column is None:
                sheet.reset_dimensions()
                sheet.calculate_dimension(force=True)
            width = max(sheet.max_column or 1, 1)
            height = sheet.max_row or 0
            if width > MAX_COLUMNS or height > MAX_ROWS or total_cells + width * height > MAX_CELLS:
                raise ValueError(f"Workbook sheet {sheet.title!r} exceeds the supported table size limit")
            name = f"sheet_{index}"
            _new_table(connection, name, width)
            slots = ", ".join("?" for _ in range(width + 1))
            sql = f'INSERT INTO "{name}" VALUES ({slots})'
            cell_rows = []
            for row_number, cells in enumerate(sheet.iter_rows(), 1):
                values = [_value(cell.value) for cell in cells[:width]]
                connection.execute(sql, (row_number, *(values + [None] * (width - len(values)))))
                for col_number, cell in enumerate(cells[:width], 1):
                    if not getattr(cell, "has_style", False):
                        continue
                    fill = cell.fill
                    font = cell.font
                    cell_rows.append((name, cell.coordinate or f"{get_column_letter(col_number)}{row_number}", row_number,
                                      col_number, _value(cell.value), fill.patternType, _color(fill.fgColor),
                                      _color(font.color), int(bool(font.bold)), cell.number_format))
            connection.executemany("INSERT INTO cell_styles VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", cell_rows)
            tables.append({"table": name, "sheet": sheet.title, "rows": height, "columns": ["row_number"] + [f"c{i}" for i in range(1, width + 1)]})
            total_cells += width * height
    finally:
        workbook.close()
    return tables


def _authorize(action: int, arg1: str | None, arg2: str | None, _db: str | None, _trigger: str | None) -> int:
    if action in {sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ}:
        return sqlite3.SQLITE_OK
    if action == sqlite3.SQLITE_FUNCTION and (arg2 or "").lower() in _SAFE_FUNCTIONS:
        return sqlite3.SQLITE_OK
    return sqlite3.SQLITE_DENY


def query_attachment(path: str | Path, sql: str) -> dict[str, object]:
    """查询 CSV/TSV/XLSX；sheet_N 保留原始行序，cell_styles 保存样式坐标。"""
    source = Path(path)
    if not source.is_file():
        raise ValueError(f"Attachment does not exist: {source}")
    if source.suffix.lower() not in {".csv", ".tsv", ".xlsx"}:
        raise ValueError("query_attachment supports CSV, TSV, and XLSX only")
    if source.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Attachment exceeds the supported file size limit")
    statement = sql.strip()
    if len(statement) > 4000 or not re.match(r"^(SELECT|WITH)\b", statement, re.IGNORECASE):
        raise ValueError("Only a SELECT query is allowed")
    connection = sqlite3.connect(":memory:")
    try:
        connection.execute("CREATE TABLE cell_styles (sheet TEXT, coordinate TEXT, row_number INTEGER, column_number INTEGER, value, fill_pattern TEXT, fill_color TEXT, font_color TEXT, bold INTEGER, number_format TEXT)")
        tables = _load_xlsx(connection, source) if source.suffix.lower() == ".xlsx" else _load_csv(connection, source)
        tables.append({"table": "cell_styles", "columns": ["sheet", "coordinate", "row_number", "column_number", "value", "fill_pattern", "fill_color", "font_color", "bold", "number_format"]})
        connection.set_authorizer(_authorize)
        steps = 0

        def progress() -> int:
            nonlocal steps
            steps += 1000
            return int(steps > MAX_SQL_STEPS)

        connection.set_progress_handler(progress, 1000)
        try:
            cursor = connection.execute(statement)
            columns = [column[0] for column in cursor.description or ()]
            result = []
            used_chars = 0
            truncated = False
            for row in cursor:
                if len(result) >= MAX_RESULT_ROWS:
                    truncated = True
                    break
                values = [_value(value) for value in row]
                used_chars += len(json.dumps(values, ensure_ascii=False, default=str))
                if used_chars > MAX_RESULT_CHARS:
                    truncated = True
                    break
                result.append(dict(zip(columns, values)))
        except sqlite3.Error as exc:
            raise ValueError(f"SQL query rejected or exceeded execution limit: {exc}") from exc
        return {"source_path": str(source.resolve()), "sha256": _fingerprint(source), "tables": tables,
                "columns": columns, "rows": result, "truncated": truncated,
                "warning": "Result truncated; narrow the query" if truncated else None}
    finally:
        connection.close()


def _decimal(value: Decimal) -> Decimal:
    if not value.is_finite() or abs(value) > MAX_MAGNITUDE:
        raise ValueError("Calculation result is non-finite or too large")
    return value


def _eval(node: ast.AST) -> Decimal:
    if isinstance(node, ast.Constant) and type(node.value) in {int, float}:
        return _decimal(Decimal(str(node.value)))
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
        operand = _eval(node.operand)
        return _decimal(operand if isinstance(node.op, ast.UAdd) else -operand)
    if isinstance(node, ast.BinOp):
        left, right = _eval(node.left), _eval(node.right)
        if isinstance(node.op, ast.Add):
            return _decimal(left + right)
        if isinstance(node.op, ast.Sub):
            return _decimal(left - right)
        if isinstance(node.op, ast.Mult):
            return _decimal(left * right)
        if isinstance(node.op, ast.Div):
            return _decimal(left / right)
        if isinstance(node.op, ast.FloorDiv):
            return _decimal(left // right)
        if isinstance(node.op, ast.Mod):
            return _decimal(left % right)
        if isinstance(node.op, ast.Pow) and right == int(right) and abs(right) <= 100:
            return _decimal(left ** int(right))
    if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and not node.keywords:
        args = [_eval(arg) for arg in node.args]
        if node.func.id == "abs" and len(args) == 1:
            return _decimal(abs(args[0]))
        if node.func.id == "sqrt" and len(args) == 1:
            return _decimal(args[0].sqrt())
        if node.func.id == "sum" and args:
            return _decimal(sum(args))
        if node.func.id == "min" and args:
            return min(args)
        if node.func.id == "max" and args:
            return max(args)
        if node.func.id == "round" and len(args) in {1, 2}:
            if len(args) == 2 and (args[1] != int(args[1]) or abs(args[1]) > 100):
                raise ValueError("round digits must be an integer from -100 to 100")
            digits = int(args[1]) if len(args) == 2 else 0
            return _decimal(round(args[0], digits))
    raise ValueError("Expression contains an unsupported operation")


def calculate(expression: str) -> dict[str, object]:
    """使用 Decimal 解释白名单算式；不调用 eval 或执行用户代码。"""
    if not expression or len(expression) > 500:
        raise ValueError("Expression must contain at most 500 characters")
    try:
        tree = ast.parse(expression, mode="eval")
        if sum(1 for _ in ast.walk(tree)) > MAX_EXPRESSION_NODES:
            raise ValueError("Expression is too complex")
        with localcontext() as context:
            context.prec = 40
            value = _eval(tree.body)
    except (SyntaxError, InvalidOperation, ArithmeticError, OverflowError) as exc:
        raise ValueError(f"Invalid calculation: {exc}") from exc
    return {"expression": expression, "result": format(value.normalize(), "f"), "truncated": False, "warning": None}
