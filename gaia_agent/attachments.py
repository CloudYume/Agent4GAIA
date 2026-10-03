"""在本地解析 GAIA 附件，并为两个模型准备相同的证据。"""

from __future__ import annotations

import base64
import math
import threading
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from PIL import Image, ImageOps
from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation

from .config import Config
from .media import VIDEO_URL, download_video, extract_video
from .sources import Task

IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
VIDEO_TYPES = {".mov", ".avi", ".mkv", ".mp4", ".webm"}
AUDIO_TYPES = {".mp3", ".mpeg", ".mpga", ".m4a", ".wav"}
TEXT_TYPES = {".txt", ".md", ".py", ".csv", ".tsv", ".json", ".xml", ".html"}
TEXT_LIMIT = 120_000
MAX_DOCUMENT_BYTES = 50 * 1024 * 1024
MAX_IMAGE_BYTES = 25 * 1024 * 1024
MAX_MEDIA_BYTES = 128 * 1024 * 1024
MAX_IMAGE_PIXELS = 36_000_000
MAX_OFFICE_EXPANDED_BYTES = 200 * 1024 * 1024
MAX_OFFICE_MEMBER_BYTES = 50 * 1024 * 1024
MAX_EMBEDDED_IMAGE_BYTES = 12 * 1024 * 1024
MAX_EMBEDDED_IMAGES = 24
MAX_PREPARED_IMAGES = 24
RESERVED_VIDEO_FRAMES = 12
MAX_SPREADSHEET_SCAN_ROWS = 100_000
_asr_lock = threading.Lock()


@dataclass
class PreparedInput:
    content: list[dict[str, Any]]
    review_text: str = ""
    review_images: list[tuple[Path, float, str]] = field(default_factory=list)
    transcript: str = ""
    warnings: list[str] = field(default_factory=list)


def _image_data_url(path: Path) -> str:
    return "data:image/jpeg;base64," + base64.b64encode(path.read_bytes()).decode("ascii")


def _normalize_image(path: Path, directory: Path, name: str) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / f"{name}.jpg"
    with Image.open(path) as image:
        if image.width * image.height > MAX_IMAGE_PIXELS:
            raise ValueError(f"Image exceeds {MAX_IMAGE_PIXELS} pixels: {path.name}")
        image = ImageOps.exif_transpose(image)
        image.thumbnail((1600, 1600))
        image.convert("RGB").save(target, format="JPEG", quality=85)
    return target


def _append_image(prepared: PreparedInput, path: Path, directory: Path, label: str, timestamp: float = 0) -> None:
    if len(prepared.review_images) >= MAX_PREPARED_IMAGES:
        prepared.warnings.append(f"{label}: omitted because the {MAX_PREPARED_IMAGES}-image budget was exhausted")
        return
    if path.stat().st_size > MAX_IMAGE_BYTES:
        raise ValueError(f"Image exceeds {MAX_IMAGE_BYTES // (1024 * 1024)} MiB limit: {path.name}")
    if path.suffix.lower() == ".gif":
        with Image.open(path) as image:
            if getattr(image, "n_frames", 1) > 1:
                prepared.warnings.append(f"{path.name}: animated image represented by its first frame")
    image = _normalize_image(path, directory, f"image-{len(prepared.review_images):03d}")
    prepared.content.append({"type": "input_text", "text": label})
    prepared.content.append({"type": "input_image", "image_url": _image_data_url(image)})
    prepared.review_images.append((image, timestamp, label))


def _append_text(prepared: PreparedInput, text: str, label: str) -> None:
    if len(text) > TEXT_LIMIT:
        prepared.warnings.append(f"{label}: text truncated at {TEXT_LIMIT} characters")
        text = text[:TEXT_LIMIT]
    payload = f"{label}:\n{text}"
    prepared.content.append({"type": "input_text", "text": payload})
    prepared.review_text += "\n" + payload


@lru_cache(maxsize=2)
def _whisper_model(name: str):
    from faster_whisper import WhisperModel

    return WhisperModel(name, device="cpu", compute_type="int8")


def _transcribe(path: Path, config: Config) -> str:
    if config.transcription_provider == "openai":
        from openai import OpenAI

        options = {"base_url": config.openai_base_url} if config.openai_base_url else {}
        client = OpenAI(api_key=config.openai_key, timeout=180, **options)
        with path.open("rb") as handle:
            result = client.audio.transcriptions.create(model=config.transcription_model, file=handle)
        if not hasattr(result, "text"):
            raise RuntimeError("Transcription API did not return a transcript")
        return result.text
    # CTranslate2 的 CPU 推理共享模型实例，串行使用可避免并发内存峰值。
    with _asr_lock:
        try:
            segments, _ = _whisper_model(config.local_asr_model).transcribe(str(path), vad_filter=True)
            lines = []
            for segment in segments:
                start = getattr(segment, "start", None)
                label = f"[{int(start // 60):02}:{int(start % 60):02}] " if start is not None else ""
                lines.append(label + segment.text.strip())
            transcript = " ".join(lines).strip()
        except (ImportError, OSError) as exc:
            raise RuntimeError(
                f"Local ASR runtime unavailable for {path.name}; reinstall the audio extra "
                "(Windows requires av==16.0.1)"
            ) from exc
        except Exception as exc:
            raise RuntimeError(f"Local ASR failed for {path.name}: {type(exc).__name__}: {exc}") from exc
    if not transcript:
        raise RuntimeError(f"No speech detected in {path.name}; this task needs audio understanding")
    return transcript


def preflight_audio(config: Config, sample_path: Path | None = None) -> dict[str, Any]:
    """Check local ASR from cached weights; optionally decode a real sample offline."""
    if config.transcription_provider != "local":
        return {"status": "api_configured", "provider": config.transcription_provider, "sample_checked": False}
    try:
        import av
        import faster_whisper  # noqa: F401
        from huggingface_hub import snapshot_download
    except (ImportError, OSError) as exc:
        return {"status": "dependency_unavailable", "reason": type(exc).__name__, "sample_checked": False}
    model_name = config.local_asr_model
    model_path = Path(model_name)
    if not model_path.is_dir():
        try:
            model_path = Path(snapshot_download(
                repo_id=f"Systran/faster-whisper-{model_name}", local_files_only=True,
                allow_patterns=["config.json", "model.bin", "tokenizer.json", "vocabulary.*"],
            ))
        except Exception:
            return {
                "status": "model_not_cached", "model": model_name,
                "av_version": getattr(av, "__version__", "unknown"), "sample_checked": False,
                "recovery": "Cache local Whisper weights or configure transcription_provider=openai",
            }
    weights = model_path / "model.bin"
    if not weights.is_file() or weights.stat().st_size == 0:
        return {
            "status": "model_not_cached", "model": model_name, "sample_checked": False,
            "recovery": "Cache local Whisper weights or configure transcription_provider=openai",
        }
    if sample_path is None:
        return {
            "status": "model_cached_unverified", "model": model_name,
            "av_version": getattr(av, "__version__", "unknown"), "sample_checked": False,
        }
    if not sample_path.is_file():
        raise FileNotFoundError(sample_path)
    try:
        with _asr_lock:
            segments, _ = _whisper_model(str(model_path)).transcribe(str(sample_path), vad_filter=True)
            transcript = " ".join(segment.text.strip() for segment in segments).strip()
    except Exception as exc:
        return {
            "status": "sample_failed", "model": model_name, "reason": type(exc).__name__,
            "av_version": getattr(av, "__version__", "unknown"), "sample_checked": True,
        }
    return {
        "status": "ready" if transcript else "no_speech_detected", "model": model_name,
        "av_version": getattr(av, "__version__", "unknown"), "sample_checked": True,
        "transcript_chars": len(transcript),
    }


def _inspect_office_archive(path: Path) -> None:
    with ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 5_000 or sum(member.file_size for member in members) > MAX_OFFICE_EXPANDED_BYTES:
            raise ValueError(f"Office archive expands beyond the safe limit: {path.name}")
        if any(member.file_size > MAX_OFFICE_MEMBER_BYTES for member in members):
            raise ValueError(f"Office archive contains an oversized member: {path.name}")


def _embedded_images(path: Path, directory: Path, prefix: str, warnings: list[str], max_images: int = MAX_EMBEDDED_IMAGES) -> list[Path]:
    images = []
    with ZipFile(path) as archive:
        names = [member for member in archive.infolist() if member.filename.startswith(prefix)]
        supported = [member for member in names if Path(member.filename).suffix.lower() in IMAGE_TYPES]
        if len(supported) < len(names):
            warnings.append(f"{path.name}: {len(names) - len(supported)} embedded media use unsupported formats")
        selected = sorted(_sample_rows(len(supported), min(max_images, MAX_EMBEDDED_IMAGES))) if max_images else []
        if len(selected) < len(supported):
            warnings.append(f"{path.name}: sampled {len(selected)} of {len(supported)} embedded images")
        for index in selected:
            member = supported[index - 1]
            if member.file_size > MAX_EMBEDDED_IMAGE_BYTES:
                warnings.append(f"{path.name}: skipped oversized embedded image {member.filename}")
                continue
            name = member.filename
            target = directory / f"embedded-{index:03d}{Path(name).suffix.lower()}"
            directory.mkdir(parents=True, exist_ok=True)
            target.write_bytes(archive.read(name))
            images.append(target)
    return images


def _sample_rows(total: int, maximum: int = 220) -> set[int]:
    if maximum < 1:
        raise ValueError("maximum must be positive")
    if total <= maximum:
        return set(range(1, total + 1))
    if maximum == 1:
        return {1}
    edge = min(20, maximum // 4)
    head = set(range(1, edge + 1))
    tail = set(range(total - edge + 1, total + 1))
    middle_count = maximum - 2 * edge
    if middle_count == 1:
        middle = {total // 2 + 1}
    else:
        first = edge + 1
        last = total - edge
        middle = {first + round(index * (last - first) / (middle_count - 1)) for index in range(middle_count)}
    return head | middle | tail


def _spreadsheet_text(path: Path) -> tuple[str, list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    lines = []
    warnings = []

    def color_text(color: Any) -> str:
        if color is None:
            return ""
        if color.type == "rgb" and isinstance(color.rgb, str):
            return f"#{color.rgb[-6:].upper()}"
        if color.type == "theme":
            return f"theme:{color.theme}"
        if color.type == "indexed":
            return f"indexed:{color.indexed}"
        return ""

    def cell_style(cell: Any) -> str:
        if not getattr(cell, "has_style", False):
            return ""
        parts = []
        fill = cell.fill
        if fill.patternType:
            colors = [color_text(fill.fgColor)]
            if fill.patternType != "solid":
                colors.append(color_text(fill.bgColor))
            parts.append(f"fill={fill.patternType}:{'/'.join(color for color in colors if color)}")
        if cell.number_format and cell.number_format != "General":
            parts.append(f"number_format={cell.number_format!r}")
        font = cell.font
        if font.name and font.name not in {"Calibri", "Aptos"}:
            parts.append(f"font={font.name}")
        if font.sz and font.sz != 11:
            parts.append(f"font_size={font.sz}")
        for label, active in (("bold", font.bold), ("italic", font.italic), ("underline", font.underline), ("strikethrough", font.strike)):
            if active:
                parts.append(label)
        if font.color is not None and not (font.color.type == "theme" and font.color.theme == 1):
            color = color_text(font.color)
            if color:
                parts.append(f"font_color={color}")
        borders = [
            side for side in ("left", "right", "top", "bottom")
            if getattr(cell.border, side) is not None and getattr(cell.border, side).style
        ]
        if borders:
            parts.append("border=" + "/".join(borders))
        if cell.alignment.horizontal:
            parts.append(f"align={cell.alignment.horizontal}")
        return "; ".join(parts)

    def render_row(index: int, row: tuple[Any, ...]) -> str:
        values = []
        styled = []
        for column, cell in enumerate(row, 1):
            value = cell.value
            values.append("" if value is None else str(value))
            style = cell_style(cell)
            if style:
                coordinate = f"{get_column_letter(column)}{index}"
                shown = "<blank>" if value is None else str(value)
                if len(shown) > 80:
                    shown = shown[:77] + "..."
                styled.append(f"{coordinate}={shown!r} [{style}]")
        text = f"Row {index}: " + "\t".join(values)
        if styled:
            text += " | Styles: " + " | ".join(styled)
        return text

    try:
        budget_per_sheet = max(0, TEXT_LIMIT // max(len(workbook.worksheets), 1) - 120)
        for sheet in workbook.worksheets:
            if not sheet.max_row:
                sheet.reset_dimensions()
                sheet.calculate_dimension(force=True)
            declared_rows = sheet.max_row or 0
            scan_rows = min(declared_rows, MAX_SPREADSHEET_SCAN_ROWS)
            sheet_lines = []
            used = 0
            full_text_fits = True
            clipped_rows = False
            for index, row in enumerate(sheet.iter_rows(values_only=False), 1):
                if index > scan_rows:
                    break
                rendered = render_row(index, row)
                if used + len(rendered) + 1 > budget_per_sheet:
                    full_text_fits = False
                    break
                sheet_lines.append(rendered)
                used += len(rendered) + 1
            if not full_text_fits:
                sample_maximum = min(220, max(2, budget_per_sheet // 100))
                selected = _sample_rows(scan_rows, sample_maximum)
                row_budget = max(0, (budget_per_sheet - len(selected)) // max(len(selected), 1))
                sheet_lines = []
                for index, row in enumerate(sheet.iter_rows(values_only=False), 1):
                    if index > scan_rows:
                        break
                    if index not in selected:
                        continue
                    rendered = render_row(index, row)
                    if len(rendered) > row_budget:
                        rendered = rendered[: max(0, row_budget - 3)] + "..." if row_budget >= 3 else rendered[:row_budget]
                        clipped_rows = True
                    sheet_lines.append(rendered)
            lines.append(f"Sheet: {sheet.title}; declared rows: {declared_rows}; shown rows: {len(sheet_lines)}")
            lines.extend(sheet_lines)
            if clipped_rows:
                warnings.append(f"{sheet.title}: some sampled rows were shortened to fit the text budget")
            if len(sheet_lines) < declared_rows:
                warnings.append(f"{sheet.title}: sampled {len(sheet_lines)} of {declared_rows} rows; the full sheet was not inspected")
            if declared_rows > scan_rows:
                warnings.append(f"{sheet.title}: rows after {MAX_SPREADSHEET_SCAN_ROWS} were not scanned")
    finally:
        workbook.close()
    return "\n".join(lines), warnings


def _document_text(path: Path) -> str:
    document = Document(path)
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _presentation_text(path: Path) -> str:
    presentation = Presentation(path)
    lines = []
    for number, slide in enumerate(presentation.slides, 1):
        lines.append(f"Slide {number}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    lines.append("\t".join(cell.text for cell in row.cells))
            if shape.has_chart:
                chart = shape.chart
                categories = []
                if chart.plots:
                    categories = [str(category.label) for category in chart.plots[0].categories]
                lines.append("Chart categories: " + ", ".join(categories))
                for series in chart.series:
                    lines.append(f"Chart series {series.name}: " + ", ".join(str(value) for value in series.values))
    return "\n".join(lines)


def _append_video(prepared: PreparedInput, path: Path, directory: Path, config: Config) -> None:
    available = max(0, MAX_PREPARED_IMAGES - len(prepared.review_images))
    frames, audio = extract_video(path, directory / "frames", max_frames=available)
    for frame, timestamp in frames:
        _append_image(prepared, frame, directory, f"Video frame near {timestamp:.1f} seconds", timestamp)
    if frames:
        prepared.warnings.append(f"{path.name}: sampled {len(frames)} video frames; brief events between frames may be missed")
    else:
        prepared.warnings.append(f"{path.name}: no video frames fit within the {MAX_PREPARED_IMAGES}-image budget")
    if audio:
        try:
            transcript = _transcribe(audio, config)
            prepared.transcript += ("\n" if prepared.transcript else "") + transcript
            _append_text(prepared, transcript, f"Video audio transcript from {path.name}")
        except Exception as exc:
            prepared.warnings.append(f"Video audio unavailable from {path.name}: {exc}")


def prepare_task(task: Task, config: Config) -> PreparedInput:
    if task.file_name and not task.attachment_path:
        raise RuntimeError(task.attachment_error or f"Missing attachment: {task.file_name}")
    prepared = PreparedInput(content=[{"type": "input_text", "text": task.question}])
    path = Path(task.attachment_path) if task.attachment_path else None
    has_video_url = bool(VIDEO_URL.search(task.question))
    if path is None and has_video_url:
        try:
            path = download_video(
                task.question, Path(".runs") / "media" / task.task_id,
                cookies_from_browser=getattr(config, "youtube_cookies_from_browser", ""),
            )
        except Exception as exc:
            raise RuntimeError(f"Cannot download video evidence: {exc}") from exc
        if path is None or not path.is_file():
            raise RuntimeError("Video URL was present but no video was downloaded")
    if path is None:
        return prepared
    if not path.is_file():
        raise FileNotFoundError(path)

    directory = Path(".runs") / "prepared" / task.task_id
    extension = path.suffix.lower()
    attachment_image_limit = MAX_PREPARED_IMAGES - RESERVED_VIDEO_FRAMES if task.attachment_path and has_video_url and extension not in VIDEO_TYPES else MAX_PREPARED_IMAGES
    limit = MAX_IMAGE_BYTES if extension in IMAGE_TYPES else MAX_MEDIA_BYTES if extension in VIDEO_TYPES | AUDIO_TYPES else MAX_DOCUMENT_BYTES
    if path.stat().st_size > limit:
        raise ValueError(f"{path.name} exceeds {limit // (1024 * 1024)} MiB limit")
    if extension in {".xlsx", ".docx", ".pptx"}:
        _inspect_office_archive(path)
    if extension in IMAGE_TYPES:
        _append_image(prepared, path, directory, f"Attachment image: {path.name}")
    elif extension in VIDEO_TYPES:
        _append_video(prepared, path, directory, config)
    elif extension in AUDIO_TYPES:
        prepared.transcript = _transcribe(path, config)
        _append_text(prepared, prepared.transcript, f"Audio transcript from {path.name}")
    elif extension == ".pdf":
        import fitz

        with fitz.open(path) as document:
            if document.page_count == 0:
                raise ValueError(f"PDF has no pages: {path.name}")
            selected = sorted(_sample_rows(document.page_count, attachment_image_limit))
            pages = []
            truncated_pages = []
            for page_number in selected:
                page = document[page_number - 1]
                page_text = page.get_text()
                if len(page_text) > 4_000:
                    truncated_pages.append(page_number)
                pages.append(f"Page {page_number}: {page_text[:4_000]}")
                target = directory / f"page-{page_number:03d}.png"
                directory.mkdir(parents=True, exist_ok=True)
                scale = min(1.25, math.sqrt(MAX_IMAGE_PIXELS / max(page.rect.width * page.rect.height, 1)))
                page.get_pixmap(matrix=fitz.Matrix(scale, scale)).save(target)
                _append_image(prepared, target, directory, f"PDF page {page_number}")
            _append_text(prepared, "\n".join(pages), path.name)
            if truncated_pages:
                prepared.warnings.append(
                    f"{path.name}: text truncated at 4000 characters on PDF pages {', '.join(map(str, truncated_pages))}"
                )
            if len(selected) < document.page_count:
                prepared.warnings.append(f"{path.name}: sampled {len(selected)} of {document.page_count} pages")
    elif extension == ".xlsx":
        spreadsheet_text, warnings = _spreadsheet_text(path)
        prepared.warnings.extend(warnings)
        _append_text(prepared, spreadsheet_text, path.name)
        for image in _embedded_images(path, directory, "xl/media/", prepared.warnings, attachment_image_limit):
            _append_image(prepared, image, directory, f"Spreadsheet embedded image: {image.name}")
    elif extension == ".docx":
        _append_text(prepared, _document_text(path), path.name)
        for image in _embedded_images(path, directory, "word/media/", prepared.warnings, attachment_image_limit):
            _append_image(prepared, image, directory, f"Document embedded image: {image.name}")
    elif extension == ".pptx":
        _append_text(prepared, _presentation_text(path), path.name)
        for image in _embedded_images(path, directory, "ppt/media/", prepared.warnings, attachment_image_limit):
            _append_image(prepared, image, directory, f"Presentation embedded image: {image.name}")
    elif extension in TEXT_TYPES:
        with path.open("rb") as handle:
            raw = handle.read(TEXT_LIMIT * 4 + 1)
        if len(raw) > TEXT_LIMIT * 4:
            prepared.warnings.append(f"{path.name}: text source exceeded read limit")
        _append_text(prepared, raw.decode("utf-8", errors="replace"), path.name)
    else:
        raise RuntimeError(f"Unsupported attachment type: {path.suffix or path.name}")
    if task.attachment_path and has_video_url and extension not in VIDEO_TYPES:
        try:
            online_video = download_video(
                task.question, Path(".runs") / "media" / task.task_id,
                cookies_from_browser=getattr(config, "youtube_cookies_from_browser", ""),
            )
        except Exception as exc:
            raise RuntimeError(f"Cannot download video evidence: {exc}") from exc
        if online_video is None or not online_video.is_file():
            raise RuntimeError("Video URL was present but no video was downloaded")
        _append_video(prepared, online_video, directory / "online-video", config)
    if prepared.warnings:
        _append_text(prepared, "\n".join(prepared.warnings), "Evidence limitations")
    return prepared
