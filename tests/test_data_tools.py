import csv

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from gaia_agent.tools.data import calculate, query_attachment


def test_query_csv_preserves_row_numbers_and_aggregates(tmp_path):
    path = tmp_path / "scores.csv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerows([["name", "score"], ["A", "12"], ["B", "30"]])
    result = query_attachment(path, "SELECT row_number, c1, c2 FROM sheet_1 WHERE CAST(c2 AS INTEGER) > 20")
    assert result["rows"] == [{"row_number": 3, "c1": "B", "c2": "30"}]
    assert result["tables"][0]["columns"] == ["row_number", "c1", "c2"]
    assert len(result["sha256"]) == 64


def test_query_xlsx_exposes_style_and_coordinate(tmp_path):
    path = tmp_path / "colors.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Sales data"
    sheet["A1"] = "amount"
    sheet["B1"] = "note"
    sheet["A2"] = 42
    sheet["A2"].fill = PatternFill(fill_type="solid", fgColor="FF00FF00")
    book.save(path)
    result = query_attachment(path, "SELECT coordinate, value, fill_color FROM cell_styles WHERE coordinate = 'A2'")
    assert result["tables"][0]["sheet"] == "Sales data"
    assert result["rows"] == [{"coordinate": "A2", "value": 42, "fill_color": "FF00FF00"}]


@pytest.mark.parametrize("sql", ["DELETE FROM sheet_1", "SELECT * FROM sheet_1; DROP TABLE sheet_1", "SELECT load_extension('x')"])
def test_query_rejects_unsafe_sql(tmp_path, sql):
    path = tmp_path / "a.csv"
    path.write_text("a\n1\n", encoding="utf-8")
    with pytest.raises(ValueError):
        query_attachment(path, sql)


def test_calculate_decimal_and_rejects_code():
    assert calculate("(0.1 + 0.2) * 100")["result"] == "30"
    assert calculate("sqrt(81) + max(1, 3)")["result"] == "12"
    with pytest.raises(ValueError):
        calculate("__import__('os').system('echo bad')")
    with pytest.raises(ValueError):
        calculate("2 ** 1000000")
    with pytest.raises(ValueError, match="round digits"):
        calculate("round(1.234, 1.5)")
