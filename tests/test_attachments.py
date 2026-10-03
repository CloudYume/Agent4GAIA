from pathlib import Path
import sys
from types import SimpleNamespace
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import fitz
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from gaia_agent import attachments
from gaia_agent.sources import Task


def _task(path: Path) -> Task:
    return Task(task_id="abc123", question="Read the attachment", level=1, file_name=path.name, attachment_path=str(path))


def test_spreadsheet_keeps_every_row_when_it_fits(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    workbook = Workbook()
    sheet = workbook.active
    long_value = "full-cell-" + "x" * 500
    for index in range(1, 301):
        sheet.append([index, long_value if index == 157 else f"value-{index}"])
    path = tmp_path / "data.xlsx"
    workbook.save(path)

    prepared = attachments.prepare_task(_task(path), SimpleNamespace())

    assert "declared rows: 300" in prepared.review_text
    assert "shown rows: 300" in prepared.review_text
    assert "Row 1: 1" in prepared.review_text
    assert "Row 157: 157" in prepared.review_text
    assert long_value in prepared.review_text
    assert "Row 300: 300" in prepared.review_text
    assert not prepared.warnings


def test_spreadsheet_reads_rows_when_dimension_metadata_is_missing(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["City", "Count"])
    sheet.append(["Paris", 17])
    sheet.append(["Lyon", 23])
    original = tmp_path / "original.xlsx"
    workbook.save(original)
    path = tmp_path / "missing-dimension.xlsx"
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(original) as source, ZipFile(path, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(data)
                dimension = root.find(f"{namespace}dimension")
                assert dimension is not None
                root.remove(dimension)
                data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(member, data)

    prepared = attachments.prepare_task(_task(path), SimpleNamespace())

    assert "declared rows: 3; shown rows: 3" in prepared.review_text
    assert "Row 2: Paris\t17" in prepared.review_text
    assert "Row 3: Lyon\t23" in prepared.review_text
    assert not prepared.warnings


def test_spreadsheet_preserves_blank_filled_cells_and_number_format(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    sheet["B1"] = 12.5
    sheet["B1"].number_format = "$#,##0.00"
    sheet["C1"] = "Header"
    sheet["C1"].font = Font(name="Arial", bold=True, color="FF00FF00")
    path = tmp_path / "styled.xlsx"
    workbook.save(path)

    prepared = attachments.prepare_task(_task(path), SimpleNamespace())

    assert "A1='<blank>' [fill=solid:#FF0000]" in prepared.review_text
    assert "B1='12.5' [number_format='$#,##0.00']" in prepared.review_text
    assert "C1='Header' [font=Arial; bold; font_color=#00FF00]" in prepared.review_text
    assert not prepared.warnings


def test_spreadsheet_style_over_budget_reports_sampling(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(attachments, "TEXT_LIMIT", 600)
    workbook = Workbook()
    sheet = workbook.active
    for index in range(1, 101):
        sheet.cell(index, 1, index).fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    path = tmp_path / "many-styles.xlsx"
    workbook.save(path)

    prepared = attachments.prepare_task(_task(path), SimpleNamespace())

    assert any("sampled" in warning for warning in prepared.warnings)
    assert "Evidence limitations" in prepared.review_text


def test_spreadsheet_samples_full_range_only_after_budget_is_exceeded(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    workbook = Workbook()
    sheet = workbook.active
    for index in range(1, 3_001):
        sheet.append([index, "x" * 200])
    path = tmp_path / "large.xlsx"
    workbook.save(path)

    prepared = attachments.prepare_task(_task(path), SimpleNamespace())

    assert "declared rows: 3000" in prepared.review_text
    assert "Row 1: 1" in prepared.review_text
    assert "Row 3000: 3000" in prepared.review_text
    assert any(f"Row {index}:" in prepared.review_text for index in range(1_000, 2_001))
    assert any("sampled" in warning for warning in prepared.warnings)
    assert "Evidence limitations" in prepared.review_text


def test_spreadsheet_scan_limit_is_explicit(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(attachments, "MAX_SPREADSHEET_SCAN_ROWS", 100)
    workbook = Workbook()
    sheet = workbook.active
    for index in range(1, 301):
        sheet.append([index])
    path = tmp_path / "bounded.xlsx"
    workbook.save(path)

    prepared = attachments.prepare_task(_task(path), SimpleNamespace())

    assert "shown rows: 100" in prepared.review_text
    assert "Row 100: 100" in prepared.review_text
    assert "Row 300: 300" not in prepared.review_text
    assert any("rows after 100 were not scanned" in warning for warning in prepared.warnings)


def test_presentation_chart_values_reach_both_models(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    chart_data = CategoryChartData()
    chart_data.categories = ["North", "South"]
    chart_data.add_series("Revenue", [12, 34])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), chart_data)
    path = tmp_path / "figures.pptx"
    presentation.save(path)

    prepared = attachments.prepare_task(_task(path), SimpleNamespace())

    assert "Chart categories: North, South" in prepared.review_text
    assert "Chart series Revenue: 12.0, 34.0" in prepared.review_text
    assert any("Chart series Revenue" in part.get("text", "") for part in prepared.content)


def test_pdf_samples_across_document_instead_of_only_first_pages(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    document = fitz.open()
    for index in range(30):
        page = document.new_page(width=120, height=120)
        page.insert_text((10, 40), f"PAGE {index + 1}")
    path = tmp_path / "report.pdf"
    document.save(path)
    document.close()

    prepared = attachments.prepare_task(_task(path), SimpleNamespace())

    assert "Page 30: PAGE 30" in prepared.review_text
    assert len(prepared.review_images) == 24
    assert prepared.review_images[-1][2] == "PDF page 30"
    assert any("sampled 24 of 30 pages" in warning for warning in prepared.warnings)


def test_pdf_page_text_truncation_is_reported(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    document = fitz.open()
    page = document.new_page(width=800, height=800)
    for index in range(90):
        page.insert_text((20, 20 + index * 8), f"LINE {index:02d} " + "x" * 64, fontsize=6)
    path = tmp_path / "long-page.pdf"
    document.save(path)
    document.close()

    with fitz.open(path) as saved:
        assert len(saved[0].get_text()) > 4_000
    prepared = attachments.prepare_task(_task(path), SimpleNamespace())

    assert len(prepared.review_images) == 1
    assert "LINE 00" in prepared.review_text
    assert any("text truncated at 4000 characters on PDF pages 1" in warning for warning in prepared.warnings)


def test_rejects_oversized_input_before_parsing(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "huge.txt"
    with path.open("wb") as handle:
        handle.truncate(attachments.MAX_DOCUMENT_BYTES + 1)

    with pytest.raises(ValueError, match="exceeds .* MiB limit"):
        attachments.prepare_task(_task(path), SimpleNamespace())


def test_rejects_oversized_office_expansion(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(attachments, "MAX_OFFICE_EXPANDED_BYTES", 128)
    path = tmp_path / "bomb.docx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * 256)

    with pytest.raises(ValueError, match="expands beyond the safe limit"):
        attachments.prepare_task(_task(path), SimpleNamespace())


def test_local_asr_joins_segments_and_rejects_silence(tmp_path, monkeypatch):
    path = tmp_path / "speech.wav"
    path.write_bytes(b"fake")

    class Model:
        def transcribe(self, path, vad_filter):
            assert vad_filter
            return iter([SimpleNamespace(text=" first ", start=0.0), SimpleNamespace(text="second", start=63.0)]), None

    monkeypatch.setattr(attachments, "_whisper_model", lambda name: Model())
    config = SimpleNamespace(transcription_provider="local", local_asr_model="small")
    assert attachments._transcribe(path, config) == "[00:00] first [01:03] second"

    class SilentModel:
        def transcribe(self, path, vad_filter):
            return iter([]), None

    monkeypatch.setattr(attachments, "_whisper_model", lambda name: SilentModel())
    with pytest.raises(RuntimeError, match="No speech detected"):
        attachments._transcribe(path, config)


def test_local_asr_dependency_failure_names_file_and_recovery(tmp_path, monkeypatch):
    path = tmp_path / "speech.wav"
    path.write_bytes(b"fake")

    def unavailable(name):
        raise OSError("application control blocked PyAV")

    monkeypatch.setattr(attachments, "_whisper_model", unavailable)
    config = SimpleNamespace(transcription_provider="local", local_asr_model="small")
    with pytest.raises(RuntimeError, match=r"speech.wav.*reinstall the audio extra.*av==16.0.1"):
        attachments._transcribe(path, config)


def test_audio_preflight_reports_missing_cached_weights_without_download(monkeypatch):
    def missing(**kwargs):
        assert kwargs["local_files_only"] is True
        raise FileNotFoundError("not cached")

    monkeypatch.setitem(sys.modules, "av", SimpleNamespace(__version__="16.0.1"))
    monkeypatch.setitem(sys.modules, "faster_whisper", SimpleNamespace())
    monkeypatch.setitem(sys.modules, "huggingface_hub", SimpleNamespace(snapshot_download=missing))
    result = attachments.preflight_audio(SimpleNamespace(transcription_provider="local", local_asr_model="small"))
    assert result["status"] == "model_not_cached"
    assert result["sample_checked"] is False


def test_audio_preflight_requires_real_successful_sample(tmp_path, monkeypatch):
    model_dir = tmp_path / "model"
    model_dir.mkdir()
    (model_dir / "model.bin").write_bytes(b"test")
    sample = tmp_path / "speech.wav"
    sample.write_bytes(b"fake sample")
    monkeypatch.setitem(sys.modules, "av", SimpleNamespace(__version__="16.0.1"))
    monkeypatch.setitem(sys.modules, "faster_whisper", SimpleNamespace())
    monkeypatch.setitem(sys.modules, "huggingface_hub", SimpleNamespace(snapshot_download=lambda **kwargs: str(model_dir)))
    monkeypatch.setattr(attachments, "_whisper_model", lambda name: SimpleNamespace(transcribe=lambda *args, **kwargs: (iter([SimpleNamespace(text="spoken words")]), None)))
    config = SimpleNamespace(transcription_provider="local", local_asr_model="small")
    assert attachments.preflight_audio(config)["status"] == "model_cached_unverified"
    result = attachments.preflight_audio(config, sample)
    assert result["status"] == "ready"
    assert result["sample_checked"] is True
    assert result["transcript_chars"] == len("spoken words")


def test_attachment_and_video_url_both_reach_solver_and_reviewer(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    text_file = tmp_path / "notes.txt"
    text_file.write_text("written evidence", encoding="utf-8")
    video_file = tmp_path / "clip.mp4"
    video_file.write_bytes(b"fake video")
    frame = tmp_path / "frame.jpg"
    Image.new("RGB", (12, 12), "red").save(frame)
    monkeypatch.setattr(attachments, "download_video", lambda question, directory, **kwargs: video_file)
    monkeypatch.setattr(attachments, "extract_video", lambda path, directory, max_frames: ([(frame, 42.0)], None))
    task = _task(text_file)
    task.question += " https://www.youtube.com/watch?v=example"

    prepared = attachments.prepare_task(task, SimpleNamespace())

    assert "written evidence" in prepared.review_text
    assert prepared.review_images[0][1:] == (42.0, "Video frame near 42.0 seconds")
    assert any("Video frame near 42.0 seconds" == part.get("text") for part in prepared.content)
    assert any("sampled 1 video frames" in warning for warning in prepared.warnings)


def test_direct_video_uses_at_most_24_images_and_warns(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    video = tmp_path / "clip.mp4"
    video.write_bytes(b"fake video")
    frame = tmp_path / "frame.jpg"
    Image.new("RGB", (12, 12), "red").save(frame)
    budgets = []

    def frames(path, directory, max_frames):
        budgets.append(max_frames)
        return [(frame, float(index)) for index in range(max_frames)], None

    monkeypatch.setattr(attachments, "extract_video", frames)
    prepared = attachments.prepare_task(_task(video), SimpleNamespace())

    assert budgets == [24]
    assert len(prepared.review_images) == 24
    assert sum(part["type"] == "input_image" for part in prepared.content) == 24
    assert any("sampled 24 video frames" in warning for warning in prepared.warnings)


def test_pdf_and_online_video_share_24_image_budget(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    document = fitz.open()
    for index in range(30):
        page = document.new_page(width=120, height=120)
        page.insert_text((10, 40), f"PAGE {index + 1}")
    pdf = tmp_path / "report.pdf"
    document.save(pdf)
    document.close()
    video = tmp_path / "clip.mp4"
    video.write_bytes(b"fake video")
    frame = tmp_path / "frame.jpg"
    Image.new("RGB", (12, 12), "red").save(frame)
    monkeypatch.setattr(attachments, "download_video", lambda question, directory, **kwargs: video)
    budgets = []

    def frames(path, directory, max_frames):
        budgets.append(max_frames)
        return [(frame, float(index)) for index in range(max_frames)], None

    monkeypatch.setattr(attachments, "extract_video", frames)
    task = _task(pdf)
    task.question += " https://www.youtube.com/watch?v=example"
    prepared = attachments.prepare_task(task, SimpleNamespace())

    assert budgets == [12]
    assert len(prepared.review_images) == 24
    assert prepared.review_images[0][2] == "PDF page 1"
    assert prepared.review_images[11][2] == "PDF page 30"
    assert prepared.review_images[12][2] == "Video frame near 0.0 seconds"
    assert any("sampled 12 of 30 pages" in warning for warning in prepared.warnings)
    assert any("sampled 12 video frames" in warning for warning in prepared.warnings)


def test_embedded_images_sample_full_archive_and_warn(tmp_path):
    path = tmp_path / "slides.pptx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for index in range(30):
            archive.writestr(f"ppt/media/image-{index:02d}.jpg", b"image")
    warnings = []
    images = attachments._embedded_images(path, tmp_path / "images", "ppt/media/", warnings, max_images=12)

    assert len(images) == 12
    assert images[0].read_bytes() == b"image"
    assert images[0].name == "embedded-001.jpg"
    assert images[-1].name == "embedded-030.jpg"
    assert any("sampled 12 of 30 embedded images" in warning for warning in warnings)
